"""Hidden Excel mirror.

SQLite is the source of truth; this module regenerates CalorieLedger.xlsx
from it (debounced via a dirty flag). The browser never touches this file —
it exists so the data is always readable/backed up as a plain workbook on
the PC. Written atomically via a temp file; if Excel has the workbook open
(Windows file lock) we skip and retry on the next flush.
"""
from __future__ import annotations

import json
import os
import threading
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .db import Database

SHEETS: list[tuple[str, str, list[str]]] = [
    ("Diary", "diary",
     ["id", "date", "time", "meal", "source", "name", "brand", "amount_mode", "amount",
      "calories", "protein", "carbs", "fat", "fibre", "accuracy", "note", "tags",
      "override", "favourite", "ref_type", "ref_id", "created_at", "updated_at", "deleted_at"]),
    ("Foods", "foods",
     ["id", "name", "brand", "barcode", "basis", "calories", "protein", "carbs", "fat",
      "fibre", "pack_g", "pack_ml", "serving_g", "items_per_pack", "serving_desc", "tags",
      "vegan", "use_count", "created_at", "updated_at", "deleted_at"]),
    ("SavedMeals", "saved_meals",
     ["id", "name", "calories", "protein", "carbs", "fat", "fibre", "components", "tags",
      "vegan", "use_count", "created_at", "updated_at", "deleted_at"]),
    ("Recipes", "recipes",
     ["id", "name", "ingredients", "cooked_yield_g", "portions", "tags", "vegan",
      "use_count", "created_at", "updated_at", "deleted_at"]),
    ("Weight", "weights", ["id", "date", "kg", "note", "created_at", "updated_at", "deleted_at"]),
    ("Waist", "waists", ["id", "date", "cm", "note", "created_at", "updated_at", "deleted_at"]),
    ("Steps", "steps", ["id", "date", "steps", "source", "created_at", "updated_at", "deleted_at"]),
    ("DayState", "day_state", ["date", "reopened", "rescues", "notified", "updated_at"]),
    ("AuditLog", "audit_log",
     ["id", "ts", "entity", "entity_id", "action", "old_json", "new_json", "op_id", "source"]),
]

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F2A44")


class ExcelStore:
    def __init__(self, db: Database, path: Path):
        self.db = db
        self.path = path
        self._dirty = threading.Event()
        self._write_lock = threading.Lock()

    def mark_dirty(self) -> None:
        self._dirty.set()

    @property
    def dirty(self) -> bool:
        return self._dirty.is_set()

    def flush_if_dirty(self) -> bool:
        if not self._dirty.is_set():
            return False
        return self.export()

    def export(self) -> bool:
        """Regenerate the workbook. Returns False if the file was locked."""
        with self._write_lock:
            wb = Workbook()
            wb.remove(wb.active)
            for sheet_name, table, cols in SHEETS:
                ws = wb.create_sheet(sheet_name)
                ws.append(cols)
                for i, _ in enumerate(cols, start=1):
                    cell = ws.cell(row=1, column=i)
                    cell.font = HEADER_FONT
                    cell.fill = HEADER_FILL
                order = "ORDER BY date" if "date" in cols else ("ORDER BY ts" if "ts" in cols else "")
                limit = "LIMIT 5000" if table == "audit_log" else ""
                if table == "audit_log":
                    order = "ORDER BY ts DESC"
                rows = self.db.query(f"SELECT * FROM {table} {order} {limit}")
                for row in rows:
                    ws.append([_cell(row.get(c)) for c in cols])
                ws.freeze_panes = "A2"
                for i, col in enumerate(cols, start=1):
                    ws.column_dimensions[get_column_letter(i)].width = max(10, min(28, len(col) + 6))

            info = wb.create_sheet("Info", 0)
            info["A1"] = "Calorie Ledger — automatic mirror of the SQLite database."
            info["A2"] = "Do not edit: this file is regenerated and changes here are overwritten."
            info["A3"] = "Deleted rows are kept with a deleted_at timestamp (soft delete)."
            info["A1"].font = Font(bold=True)

            tmp = self.path.with_suffix(".xlsx.tmp")
            try:
                wb.save(tmp)
                os.replace(tmp, self.path)
            except PermissionError:
                # Workbook open in Excel on Windows — try again on next flush.
                tmp.unlink(missing_ok=True)
                return False
            self._dirty.clear()
            return True


def _cell(value):
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return value
