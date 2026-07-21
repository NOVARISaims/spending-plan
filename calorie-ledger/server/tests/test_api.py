import uuid

from conftest import PASSPHRASE, sync_op


def test_health_and_auth_flow(client):
    r = client.get("/api/health")
    assert r.status_code == 200 and r.json()["auth_configured"] is True

    assert client.get("/api/today").status_code == 401
    r = client.post("/api/auth/login", json={"passphrase": "wrong-wrong", "label": "x"})
    assert r.status_code == 401

    r = client.post("/api/auth/login", json={"passphrase": PASSPHRASE, "label": "iphone"})
    token = r.json()["token"]
    r = client.get("/api/today", headers={"Authorization": f"Bearer {token}"})
    assert r.status_code == 200

    r = client.post("/api/auth/logout", headers={"Authorization": f"Bearer {token}"})
    assert r.status_code == 200
    r = client.get("/api/today", headers={"Authorization": f"Bearer {token}"})
    assert r.status_code == 401


def test_food_lifecycle_and_barcode(auth_client):
    res = sync_op(auth_client, "food", "create", {
        "name": "Peanut Butter", "brand": "Nutty", "barcode": "5012345678900",
        "basis": "per_100g", "calories": 600, "protein": 25, "fat": 50,
        "pack_g": 340, "serving_g": 15, "tags": ["protein"], "vegan": True,
    })
    assert res["status"] == "applied"
    food_id = res["result"]["id"]

    r = auth_client.get("/api/catalog/foods", params={"barcode": "5012345678900"})
    assert r.json()["matched"] is True
    assert r.json()["foods"][0]["id"] == food_id

    r = auth_client.get("/api/catalog/foods", params={"barcode": "0000000000000"})
    assert r.json()["matched"] is False

    res = sync_op(auth_client, "food", "delete", {"id": food_id})
    assert res["status"] == "applied"
    r = auth_client.get("/api/catalog/foods", params={"barcode": "5012345678900"})
    assert r.json()["matched"] is False  # soft-deleted rows are hidden

    r = auth_client.get("/api/audit", params={"entity": "food", "entity_id": food_id})
    actions = [a["action"] for a in r.json()["audit"]]
    assert "create" in actions and "soft_delete" in actions


def test_diary_create_computes_from_food(auth_client):
    res = sync_op(auth_client, "food", "create", {
        "name": "Oats", "basis": "per_100g", "calories": 380, "protein": 13,
        "fibre": 10, "pack_g": 500, "tags": ["fibre", "wholegrain"], "vegan": True,
    })
    food_id = res["result"]["id"]
    res = sync_op(auth_client, "diary", "create", {
        "date": "2026-07-20", "time": "08:00", "meal": "breakfast",
        "source": "barcode_product", "ref_type": "food", "ref_id": food_id,
        "name": "Oats", "amount_mode": "grams", "amount": 60, "calories": 0,
    })
    assert res["status"] == "applied"
    assert res["result"]["calories"] == 228  # server-computed, not the client's 0

    r = auth_client.get("/api/diary", params={"date": "2026-07-20"})
    entry = r.json()["entries"][0]
    assert entry["fibre"] == 6.0
    assert entry["tags"] == ["fibre", "wholegrain"]  # inherited from food

    # last-used amount mode remembered on the food
    r = auth_client.get(f"/api/catalog/foods/{food_id}")
    assert r.json()["last_amount_mode"] == "grams"
    assert r.json()["use_count"] == 1


def test_sync_idempotency(auth_client):
    op_id = str(uuid.uuid4())
    payload = {"date": "2026-07-20", "time": "13:00", "meal": "lunch",
               "source": "quick", "name": "Sandwich", "calories": 450,
               "accuracy": "good_estimate"}
    first = sync_op(auth_client, "diary", "create", payload, op_id=op_id)
    dup = sync_op(auth_client, "diary", "create", payload, op_id=op_id)
    assert first["status"] == "applied"
    assert dup["status"] == "duplicate"
    r = auth_client.get("/api/diary", params={"date": "2026-07-20"})
    assert len([e for e in r.json()["entries"] if e["name"] == "Sandwich"]) == 1


def test_diary_edit_soft_delete_undelete(auth_client):
    res = sync_op(auth_client, "diary", "create", {
        "date": "2026-07-19", "time": "19:00", "meal": "dinner",
        "source": "quick", "name": "Curry", "calories": 700,
    })
    entry_id = res["result"]["id"]
    res = sync_op(auth_client, "diary", "update",
                  {"id": entry_id, "calories": 650, "meal": "snack", "note": "smaller portion"})
    assert res["status"] == "applied" and res["result"]["calories"] == 650

    res = sync_op(auth_client, "diary", "delete", {"id": entry_id})
    assert res["status"] == "applied"
    r = auth_client.get("/api/diary", params={"date": "2026-07-19"})
    assert r.json()["entries"] == []

    res = sync_op(auth_client, "diary", "undelete", {"id": entry_id})
    assert res["status"] == "applied"
    r = auth_client.get("/api/diary", params={"date": "2026-07-19"})
    assert r.json()["entries"][0]["note"] == "smaller portion"

    r = auth_client.get("/api/audit", params={"entity": "diary", "entity_id": entry_id})
    actions = [a["action"] for a in r.json()["audit"]]
    assert {"create", "update", "soft_delete", "undelete"} <= set(actions)


def test_food_edit_propagates_with_audit(auth_client):
    res = sync_op(auth_client, "food", "create", {
        "name": "Yogurt", "basis": "per_100g", "calories": 60, "pack_g": 500,
    })
    food_id = res["result"]["id"]
    for day in ("2026-07-10", "2026-07-11"):
        sync_op(auth_client, "diary", "create", {
            "date": day, "time": "08:00", "meal": "breakfast", "source": "manual_product",
            "ref_type": "food", "ref_id": food_id, "name": "Yogurt",
            "amount_mode": "grams", "amount": 200, "calories": 0,
        })

    r = auth_client.get(f"/api/catalog/foods/{food_id}/impact")
    assert r.json()["linked_entries"] == 2

    # propagate=none: history untouched
    res = sync_op(auth_client, "food", "update", {
        "id": food_id, "propagate": "none", "name": "Yogurt",
        "basis": "per_100g", "calories": 90, "pack_g": 500,
    })
    assert res["result"]["recalculated_entries"] == 0
    r = auth_client.get("/api/diary", params={"date": "2026-07-10"})
    assert r.json()["entries"][0]["calories"] == 120  # still 60/100g * 200g

    # propagate=past: entries recalculated, old values preserved in audit
    res = sync_op(auth_client, "food", "update", {
        "id": food_id, "propagate": "past", "name": "Yogurt",
        "basis": "per_100g", "calories": 90, "pack_g": 500,
    })
    assert res["result"]["recalculated_entries"] == 2
    r = auth_client.get("/api/diary", params={"date": "2026-07-10"})
    assert r.json()["entries"][0]["calories"] == 180

    r = auth_client.get("/api/audit", params={"entity": "diary"})
    recalcs = [a for a in r.json()["audit"] if a["action"] == "recalc"]
    assert len(recalcs) == 2
    assert '"calories": 120' in recalcs[0]["old_json"]


def test_body_upserts_and_today(auth_client):
    import datetime
    today = datetime.date.today().isoformat()
    res = sync_op(auth_client, "weight", "log", {"date": today, "kg": 82.4})
    assert res["status"] == "applied"
    res = sync_op(auth_client, "weight", "log", {"date": today, "kg": 82.1})
    assert res["status"] == "applied"
    r = auth_client.get("/api/body")
    assert len([w for w in r.json()["weights"] if w["date"] == today]) == 1
    assert r.json()["latest_weight"]["kg"] == 82.1

    sync_op(auth_client, "waist", "log", {"date": today, "cm": 91.5})
    sync_op(auth_client, "steps", "log", {"date": today, "steps": 8500})
    sync_op(auth_client, "diary", "create", {
        "date": today, "time": "08:00", "meal": "breakfast",
        "source": "quick", "name": "Toast", "calories": 300,
    })

    r = auth_client.get("/api/today")
    t = r.json()
    assert t["eaten"] == 300
    assert t["left"] == t["target"] - 300
    assert t["weight"]["latest"]["kg"] == 82.1
    assert t["waist"]["latest"]["cm"] == 91.5
    assert t["steps"] == 8500
    assert t["meal_totals"]["breakfast"] == 300
    assert t["weekly"]["eaten"] >= 300


def test_day_reopen_and_rescue(auth_client, app):
    """Rescue moves expired remainder into a later window."""
    from app.windows import day_windows
    services = app.state.services

    today = __import__("datetime").date.today().isoformat()
    res = sync_op(auth_client, "day", "reopen", {"date": today, "window_id": "breakfast"})
    assert res["status"] == "applied"

    # Rescue validity depends on wall-clock; exercise engine directly for determinism.
    settings = services.settings()
    day = {"reopened": [], "rescues": [{"from": "breakfast", "to": "dinner", "amount": 300}]}
    pic = day_windows(settings, {}, day, 18 * 60)
    dinner = next(w for w in pic["windows"] if w["id"] == "dinner")
    assert dinner["allowance"] == settings["windows"][2]["allowance"] + 300

    res = sync_op(auth_client, "day", "rescue",
                  {"date": today, "from_window": "breakfast", "to_window": "breakfast"})
    assert res["status"] == "error"


def test_settings_update_and_validation(auth_client):
    res = sync_op(auth_client, "settings", "update", {"daily_target": 2000})
    assert res["status"] == "applied"
    r = auth_client.get("/api/settings")
    assert r.json()["daily_target"] == 2000

    res = sync_op(auth_client, "settings", "update", {"daily_target": 100})
    assert res["status"] == "error"
    res = sync_op(auth_client, "settings", "update", {"windows": [
        {"id": "breakfast", "start": "09:00", "end": "08:00", "allowance": 400, "enabled": True}
    ]})
    assert res["status"] == "error"
    res = sync_op(auth_client, "settings", "update", {"nonsense_key": 1})
    assert res["status"] == "error"


def test_validation_rejects_junk(auth_client):
    res = sync_op(auth_client, "diary", "create", {
        "date": "21/07/2026", "time": "08:00", "meal": "breakfast",
        "source": "quick", "name": "X", "calories": 100,
    })
    assert res["status"] == "error"
    res = sync_op(auth_client, "weight", "log", {"date": "2026-07-20", "kg": 1000})
    assert res["status"] == "error"
    res = sync_op(auth_client, "diary", "create", {
        "date": "2026-07-20", "time": "08:00", "meal": "elevenses",
        "source": "quick", "name": "X", "calories": 100,
    })
    assert res["status"] == "error"


def test_excel_backup_journal(auth_client, app):
    sync_op(auth_client, "diary", "create", {
        "date": "2026-07-20", "time": "12:30", "meal": "lunch",
        "source": "restaurant", "name": "Burrito", "calories": 800,
        "accuracy": "rough_estimate",
    })
    r = auth_client.post("/api/export/excel")
    assert r.status_code == 200
    assert app.state.cfg.excel_path.exists()

    from openpyxl import load_workbook
    wb = load_workbook(app.state.cfg.excel_path)
    assert {"Diary", "Foods", "Weight", "AuditLog", "Info"} <= set(wb.sheetnames)
    diary_rows = list(wb["Diary"].values)
    assert any("Burrito" in str(row) for row in diary_rows[1:])

    r = auth_client.post("/api/backups/run")
    assert r.status_code == 200 and r.json()["file"].startswith("ledger-")
    r = auth_client.get("/api/backups")
    assert len(r.json()["backups"]) == 1

    journal = app.state.cfg.journal_path.read_text().strip().splitlines()
    assert any('"entity": "diary"' in line for line in journal)


def test_shortcut_steps_endpoint(auth_client):
    r = auth_client.post("/api/shortcuts/steps", json={"steps": 12000})
    assert r.status_code == 200 and r.json()["status"] == "applied"
    r = auth_client.get("/api/today")
    assert r.json()["steps"] == 12000


def test_saved_meal_and_recipe_flow(auth_client):
    res = sync_op(auth_client, "saved_meal", "create", {
        "name": "Protein bowl", "calories": 520, "protein": 35, "fibre": 12,
        "components": [{"name": "tofu", "calories": 180}, {"name": "rice", "calories": 340}],
        "tags": ["protein", "fibre"], "vegan": True,
    })
    meal_id = res["result"]["id"]
    res = sync_op(auth_client, "diary", "create", {
        "date": "2026-07-20", "time": "13:00", "meal": "lunch", "source": "saved_meal",
        "ref_type": "saved_meal", "ref_id": meal_id, "name": "Protein bowl",
        "amount_mode": "servings", "amount": 0.5, "calories": 0,
    })
    assert res["result"]["calories"] == 260

    res = sync_op(auth_client, "recipe", "create", {
        "name": "Dal", "ingredients": [
            {"name": "lentils", "amount_g": 300, "calories": 345, "protein": 27, "fibre": 24},
            {"name": "onion", "amount_g": 150, "calories": 60, "fibre": 3},
        ],
        "cooked_yield_g": 900, "portions": 3, "vegan": True,
    })
    recipe_id = res["result"]["id"]
    assert res["result"]["totals"]["per_portion"]["calories"] == 135
    res = sync_op(auth_client, "diary", "create", {
        "date": "2026-07-20", "time": "19:00", "meal": "dinner", "source": "recipe",
        "ref_type": "recipe", "ref_id": recipe_id, "name": "Dal",
        "amount_mode": "portion", "amount": 2, "calories": 0,
    })
    assert res["result"]["calories"] == 270

    # editing the recipe with propagate=past recalculates the entry
    res = sync_op(auth_client, "recipe", "update", {
        "id": recipe_id, "propagate": "past", "name": "Dal", "ingredients": [
            {"name": "lentils", "amount_g": 300, "calories": 345, "protein": 27, "fibre": 24},
            {"name": "onion", "amount_g": 150, "calories": 60, "fibre": 3},
            {"name": "oil", "amount_g": 30, "calories": 270},
        ],
        "cooked_yield_g": 900, "portions": 3, "vegan": True,
    })
    assert res["result"]["recalculated_entries"] == 1
    r = auth_client.get("/api/diary", params={"date": "2026-07-20"})
    dal = next(e for e in r.json()["entries"] if e["name"] == "Dal")
    assert dal["calories"] == 450


def test_progress_and_review(auth_client):
    import datetime
    monday = datetime.date.today() - datetime.timedelta(days=datetime.date.today().weekday() + 7)
    for i in range(7):
        d = (monday + datetime.timedelta(days=i)).isoformat()
        sync_op(auth_client, "diary", "create", {
            "date": d, "time": "12:30", "meal": "lunch", "source": "quick",
            "name": f"Meal {i}", "calories": 1900,
        })
    r = auth_client.get("/api/progress", params={"days": 30})
    assert r.status_code == 200
    body = r.json()
    assert len(body["calories"]) >= 7
    assert body["weekly"]

    r = auth_client.get("/api/review/weekly")
    review = r.json()
    assert review["days_logged"] == 7
    assert review["eaten"] == 1900 * 7
    assert isinstance(review["suggestion"], str) and len(review["suggestion"]) > 10
